"""
Extracts plain text from DOCX bytes using python-docx.
Walks paragraphs and table cells, since resumes sometimes use tables for layout.
"""

import io
from docx import Document


def extract_text_from_docx(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))
    chunks = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            chunks.append(paragraph.text.strip())

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    chunks.append(cell.text.strip())

    return "\n".join(chunks)